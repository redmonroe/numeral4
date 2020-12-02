from datetime import datetime
from config import Config
import os
import os.path
from pathlib import Path
from openpyxl import Workbook
from openpyxl import load_workbook
from decimal import Decimal


def sheet_finder(function, path):
    print('\nHelping you find a sheet!!!!!\n')
    print(f'for {function}.')

    file_path = find_last_modified(path)
    print(f'\nThis is most recent active file: {file_path.name}\n\n')

    choice = int(input('If you want to use most recent active file PRESS 1 \n or PRESS another key to keep searching'))

    if choice == 1:
        return file_path
    else:
        path = walk_download_folder(path)
        return path


class Liltilities(object):

    def __init__(self):
        self.heck = None

    def hello(self):
        print('hello')

    def get_book_name(self, service, sh_id):
            response = service.spreadsheets().get(
                spreadsheetId=sh_id
                ).execute()

            book_name = response['properties']['title']
            return book_name

    def get_existing_sheets(self, service, sh_id):
        response = service.spreadsheets().get(
            spreadsheetId=sh_id
            ).execute()

        # print(response.keys())
        print("book name:", response['properties']['title'])
        print('\n')
        print("URL:", response['spreadsheetUrl'])
        print('\n')
        print("Sheet name ::::::: sheet id ")
        print("********** ::::::: *********")

        titles = []
        sheet_ids = []
        for i in range(len(response['sheets'])):
            title = (response['sheets'][i]['properties']['title'])
            sheet_id = (response['sheets'][i]['properties']['sheetId'])
            titles.append(title)
            sheet_ids.append(sheet_id)
        titles_dict = dict(zip(titles, sheet_ids))

        return titles_dict

    def existing_ids(self, dict):
        # creates index of titles and ids
        idx = []
        sheet = []
        id = []
        sub = ()
        idx_dict = {}
        print("======================================")
        print("\n\n")
        for k, v in enumerate(dict):
            idx.append(k)
            sheet.append(v)
            id.append(dict[v])
            print(k, "***", v, "***", dict[v])

        print("\n")
        print("======================================")
        sub = tuple(zip(sheet, id))
        idx_list = list(zip(idx, sub))

        return idx_list

    def show_files_as_choices(self, list):
        choice = []
        files = []
        choice_file = {}

        for count, (k, v) in enumerate(list.items(), 0):
            print(count, "****", k, '****', v)
            choice.append(count)
            files.append(k)

        selection = int(input("Please select an item to work with:"))

        choice_file = dict(zip(choice, files))

        for k, v in choice_file.items():
            if selection == k:
                return v, selection

    def get_letter_by_choice(self, choice, offset):
        print(choice)
        if isinstance(choice, list):
            print('this is type list.')
            for item in choice:
                choice = item
                print(choice)

        letters = 'abcdefghijklm'
        choice = choice - offset
        output = letters[choice]

        return output

    def string_to_decimal(self, input):

        input = input.replace(',', '')
        input = input.replace('$', '')

        checking = isinstance(input, str)
        if checking == True:
            print('Found a string . . . converting to decimal.')
            output = Decimal(input)
        else:
            print(type(input))

        return output

    def string_slicer(self, input, start_slice, end_slice):
        checking = isinstance(input, str)
        if checking == True:
            output = input[start_slice:end_slice]

        return output

    def non_secure_random_number(self):
        import random
        random.seed(version=2)
        random_id = random.randrange(100000)
        return random_id

    def sheet_finder(self, use_case, path):
        print('\nHelping you find a sheet!!!!!\n')
        print(f'for {use_case}.')

        time, file_path = max((file.stat().st_mtime, file) for file in path.iterdir())
        print(f'\nThis is most recent active file: {file_path.name}\n\n')

        choice = int(input('If you want to use most recent active file PRESS 1 \n or PRESS another key to keep searching'))

        if choice == 1:
            return file_path
        else:
            path = self.walk_download_folder(path)
            return path

    def find_last_modified(self, dir):
        time, file_path = max((file.stat().st_mtime, file) for file in path.iterdir())
        # print(datetime.fromtimestamp(time), file_path.name)
        return file_path

    def walk_download_folder(self, path):
        choice = []
        files = []
        choice_file = {}

        # path = ''

        for count, file in enumerate(path.iterdir()):
            print('\n', count, "******", file.name)
            choice.append(count)
            files.append(file)

        selection = int(input("Please select an item to work with:"))

        choice_file = dict(zip(choice, files))

        for k, v in choice_file.items():
            if selection == k:
                path = v
        return path

    def load_activate_workbook(self, path):
        workbook = load_workbook(path)
        sheet = workbook.active
        print(f"Opening {sheet} from book {workbook} . . .  ")
        return sheet

    def generic_db_revision(self):
        os.system('alembic stamp head')
        os.system('alembic revision --autogenerate -m "upgraded from commandline"')
        os.system('alembic upgrade head')
        print('ok')

class DB_Utilities(object):

# https://www.postgresql.org/docs/8.0/backup.html
# notes on backup and restore
    @staticmethod
    def pg_dump_one():
        from datetime import datetime as dt
        bu_time = dt.now()
        print(bu_time)
        # set PGPASSFILE = C:\Users\joewa\AppData\Roaming\postgresql\pgpass.conf
        os.system(f'pg_dump --dbname={Config.PG_DUMPS_URI} > "{Config.DB_BACKUPS}\dump{bu_time.month}{bu_time.day}{bu_time.year}{bu_time.hour}.sql"')


    @staticmethod
    def pg_restore_one(infile, testing=True):
        print('infile name:', infile)
        os.system(f'pg_dump --dbname={xxxx} > "{infile}')
    # psql dbname < infile

    @staticmethod
    def backup_db():
        import gzip
        from sh import pg_dump
        with gzip.open('backup.gz', 'wb') as f:
            pg_dump('-h', 'localhost', '-U', 'postgres', 'my_dabatase', _out=f)
    
    @staticmethod
    def show_existing_tables(engine):
        print("DB URI:", engine)
        for item in enumerate(engine.table_names()):
            print(item)

    @staticmethod
    def recreate_database(engine, Base):
        print("Dropping all tables and recreating . . . ")
        print("DB URI:", engine)
        Base.metadata.drop_all(engine)
        Base.metadata.create_all(engine)

    @staticmethod
    def recreate_no_drop(engine, Base, Session):
        """I can delete a table at pg admin, then use this functions
        to recreate schema, without overwriting/losing other tables"""
        s = Session()
        Base.metadata.create_all(engine, Base.metadata.tables.values(),checkfirst=True)

    @staticmethod
    def show_class_structure(clsname): # this is how you import a class from a string
        omar = sys.modules[__name__]
        instance = getattr(omar, clsname)
        print(dir(instance))
